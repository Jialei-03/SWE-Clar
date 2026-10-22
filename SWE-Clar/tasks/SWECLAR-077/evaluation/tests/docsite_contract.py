import os,json,base64,select,tempfile,subprocess,threading,uuid,importlib.util,sys,time,traceback,urllib.request,shutil,zipfile
from pathlib import Path
from urllib.parse import urljoin
from copy import deepcopy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from io import BytesIO
from playwright.sync_api import sync_playwright
sys.path.insert(0,"/public_tests")
from browser_isolation import launch
from fixtures import assets
from pypdf import PdfReader,PdfWriter
from PIL import Image
ROOT=Path(os.environ.get('SWECLAR_PROJECT_DIR','.')).resolve();TESTS=Path(__file__).resolve().parent


class App:
 def __init__(self,work):
  self.work=work;work.chmod(0o777);self.source=work/'source';self.out=work/'output';self.source.mkdir();self.out.mkdir();self.out.chmod(0o777)
  self.lines=['First ordered paragraph','Repeated line','Repeated line','Final page paragraph'];assets(self.source,self.lines);shutil.copy2(self.source/'input.pdf',work/'outside.pdf');shutil.copy2(self.source/'image.png',work/'outside.png')
  env={'PATH':os.environ['PATH'],'HOME':'/tmp','CLARSWE_PRIVATE_SERVICE_SENTINEL':'private-parent-only'}
  self.worker=subprocess.Popen([sys.executable,'-I','/public_tests/worker.py',str(ROOT)],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,env=env)
  assert select.select([self.worker.stdout],[],[],5)[0];ready=json.loads(self.worker.stdout.readline());assert ready['uid']==1000 and ready['cap_eff']==0 and set(ready['environment_keys'])<={'HOME','PATH'}
  self.php=subprocess.Popen([sys.executable,'-I','/public_tests/php_launcher.py',str(ROOT),str(self.out),str(work/'php-proof.json')],stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,env=env)
  assert select.select([self.php.stdout],[],[],5)[0];port=json.loads(self.php.stdout.readline())['port'];self.url='http://127.0.0.1:'+str(port)+'/index.php?page='
  proof=json.loads((work/'php-proof.json').read_text());status=dict(x.split(':',1) for x in Path(f'/proc/{self.php.pid}/status').read_text().splitlines() if ':' in x);assert set(status['Uid'].split())=={'1000'} and int(status['CapEff'],16)==0 and proof['parent_sentinel_absent']
  Path('/private/results').mkdir(parents=True,exist_ok=True);open('/private/results/service-isolation.jsonl','a').write(json.dumps(proof)+'\n');open('/private/results/worker-isolation.jsonl','a').write(json.dumps(ready)+'\n')
 def close(self):
  for p in [self.worker,self.php]:
   if p.poll() is None:p.terminate();p.wait(timeout=5)
 def call(self,s,a):
  before=json.dumps([s,a],sort_keys=True);self.worker.stdin.write(json.dumps({'state':s,'action':a,'output_root':str(self.out),'source_root':str(self.source)})+'\n');self.worker.stdin.flush();assert select.select([self.worker.stdout],[],[],10)[0],'candidate worker timeout';v=json.loads(self.worker.stdout.readline());assert 'exception' not in v,v;assert json.dumps([v['state'],v['action']],sort_keys=True)==before and isinstance(v['result'],dict) and not v['network_attempts'];return v['result']
 def read(self,path):
  r=subprocess.run([sys.executable,'-I','/public_tests/files.py',str(path)],env={'PATH':os.environ['PATH'],'HOME':'/tmp','CLARSWE_PRIVATE_SERVICE_SENTINEL':'private-parent-only'},capture_output=True,text=True,timeout=5);assert r.returncode==0,r.stderr;return base64.b64decode(r.stdout)
 def probe(self,pid):
  with sync_playwright() as p:
   browser=launch(p.chromium,headless=True)
   try:
    page=browser.new_page();page.goto(self.url+pid,wait_until='networkidle')
    result=page.evaluate("()=>({title:document.title,heading:document.querySelector('h1').textContent,body:document.querySelector('.intro').textContent,sections:[...document.querySelectorAll('section h2')].map(x=>x.textContent),images:[...document.images].map(x=>({width:x.naturalWidth,height:x.naturalHeight})),injected:!!window.injected,description:document.querySelector('meta[name=description]')?.content,canonical:document.querySelector('link[rel=canonical]')?.getAttribute('href'),robots:document.querySelector('meta[name=robots]')?.content})");result['layouts']=[]
    result['image_pixels']=page.locator('img').evaluate_all('(xs)=>xs.map(x=>{const c=document.createElement("canvas");c.width=x.naturalWidth;c.height=x.naturalHeight;const g=c.getContext("2d");g.drawImage(x,0,0);return Array.from(g.getImageData(0,0,c.width,c.height).data)})')
    for width in [360,599,600,1023,1024,1440]:
     page.set_viewport_size({'width':width,'height':800});result['layouts'].append(page.evaluate("()=>{const boxes=[...document.querySelectorAll('.sections section')].map(x=>x.getBoundingClientRect());return {width:innerWidth,columns:boxes.filter(b=>Math.abs(b.y-boxes[0].y)<2).length,overflow:document.documentElement.scrollWidth>innerWidth}}"))
    return result
   finally:browser.close()

def base():return {'assets':{'sample':{'path':'image.png'}},'pdfs':{'input':{'type':'pdf','path':'input.pdf'}},'letterheads':{'plain':{'path':'letterhead.docx'}}}
def page(pid='home'):return {'type':'upsert_page','page_id':pid,'title':'Example course','body':'Literal <script>window.injected=true</script> content','sections':['Introduction','Lessons','Exercises','Reading','Practice','Review'],'asset_ids':['sample'],'viewport_width':500}
def render(a,pid='home'):
 s=a.call(base(),page(pid));assert 'error' not in s,s;return s

def assert_image_pixels(actual,expected):
 assert actual['image_pixels']==[list(expected)],'Displayed image differs from registered PNG pixels'

def ru001(a):
 state=base();request=page();first=a.call(state,request);assert 'error' not in first and 'error' not in a.call(state,request)
 assert (a.out/first['result']['file']).is_file();assert a.call(state,{'type':'unknown'})['error']['code']=='unknown_action'

def ru002(a):
 # Different patterns with identical dimensions distinguish actual source use
 # from a same-size placeholder, without constraining PNG encoding or URLs.
 for index in range(2):
  im=Image.new('RGB',(120,80));im.putdata([((x*7+index*83)%256,(y*11+index*47)%256,((x+y)*3+index*91)%256) for y in range(80) for x in range(120)])
  im.save(a.source/f'pattern-{index}.png');state=base();state['assets']['sample']['path']=f'pattern-{index}.png'
  assert 'error' not in a.call(state,page(f'pattern-{index}'))
  actual=a.probe(f'pattern-{index}')
  assert_image_pixels(actual,im.convert('RGBA').tobytes())
 s=render(a,'responsive');r=a.probe('responsive')
 assert r['title']=='Example course' and r['heading']=='Example course' and r['body']==page()['body']
 assert r['sections']==page()['sections'] and r['images']==[{'width':120,'height':80}],r
 assert not r['injected']
 for row in r['layouts']:
  expected=1 if row['width']<600 else 2 if row['width']<1024 else 3
  assert row['columns']==expected and not row['overflow'],row
 s['pages']['responsive']['custom_metadata']={'category':'synthetic'};update=page('responsive');update['title']='Updated course';update['sections']=['Only new section'];s=a.call(s,update);r=a.probe('responsive');assert r['heading']=='Updated course' and r['sections']==['Only new section'] and s['pages']['responsive']['custom_metadata']=={'category':'synthetic'}

def ru003(a):
 s=render(a,'seo');s=a.call(s,{'type':'set_seo','page_id':'seo','title':'  Search title  ','description':'  Learn sample material  ','canonical_path':'/resources/sample'})
 assert all(s['result'][k]==v for k,v in {'title':'Search title','description':'Learn sample material','canonical_path':'/resources/sample','indexable':True}.items())
 r=a.probe('seo');assert r['title']=='Search title' and r['description']=='Learn sample material' and urljoin(a.url,r['canonical'])==urljoin(a.url,'/resources/sample') and {'index','follow'}<=set(x.strip().lower() for x in r['robots'].split(',')) and not {'noindex','nofollow'}&set(x.strip().lower() for x in r['robots'].split(','))
 assert r['body']==page()['body'] and r['sections']==page()['sections']
 update=page('seo');update['title']='New body heading';s=a.call(s,update);r=a.probe('seo');assert r['title']=='Search title' and r['heading']=='New body heading'
 for path in ['https://example.test/path','//example.test/path','/path\\escape','/path\nnext']:
  before=deepcopy(s);r=a.call(s,{'type':'set_seo','page_id':'seo','title':'T','description':'D','canonical_path':path});assert r['error']['code']=='invalid_seo';r.pop('error');assert all(r.get(k)==v for k,v in before.items() if k!='result')

def ru004(a):
 s=a.call(base(),{'type':'convert_pdf','document_id':'converted','pdf_id':'input','letterhead_id':'plain'});assert 'error' not in s,s
 f=a.out/s['result']['file'];assert f.is_file() and s['result']['editable'] is True
 d=Document(BytesIO(a.read(f)));assert [p.text for p in d.paragraphs]==['Sample resource note',*a.lines]
 assert [p.text for p in d.sections[0].header.paragraphs]==['SYNTHETIC LETTERHEAD']
 assert not d.core_properties.author and not d.core_properties.last_modified_by
 assert len(d.inline_shapes)==0
 d.paragraphs[-1].text='Edited after conversion';d.save(a.out/'edited.docx');assert Document(a.out/'edited.docx').paragraphs[-1].text=='Edited after conversion'
 # A second real source defeats fixed fixture text and supplied fake blocks.
 more=a.source/'more';lines=['Different PDF first line','Second PDF line','Second PDF line','Different final line'];assets(more,lines);state=base();state['pdfs']['input']['path']='more/input.pdf';state['pdfs']['input']['blocks']=['Invented block'];other=a.call(state,{'type':'convert_pdf','document_id':'second_doc','pdf_id':'input','letterhead_id':'plain'});assert [p.text for p in Document(BytesIO(a.read(a.out/other['result']['file']))).paragraphs]==['Sample resource note',*lines]
 # Existing output and state IDs cannot be silently overwritten.
 before=a.read(f);r=a.call(s,{'type':'convert_pdf','document_id':'converted','pdf_id':'input','letterhead_id':'plain'});assert r['error']['code']=='invalid_document' and a.read(f)==before
 r=a.call(base(),{'type':'convert_pdf','document_id':'converted','pdf_id':'input','letterhead_id':'plain'});assert r['error']['code']=='invalid_document' and a.read(f)==before


def ru005(a):
 s=render(a,'stable');old=a.read(a.out/'stable.json')
 for value in [False,0,-1,float('inf'),float('nan')]:
  request=page('stable');request['viewport_width']=value;before=deepcopy(s);r=a.call(s,request);assert r['error']['code']=='invalid_page' and a.read(a.out/'stable.json')==old
 request=page('stable');request['asset_ids']=['unknown'];assert a.call(s,request)['error']['code']=='invalid_page'
 assert a.call(s,{'type':'set_seo','page_id':'missing','title':'T','description':'D','canonical_path':'/x'})['error']['code']=='invalid_seo'
 for ident,path,kind in [('blank','blank.pdf','pdf'),('locked','encrypted.pdf','pdf'),('wrong','image.png','png'),('absent','absent.pdf','pdf'),('escape','../outside.pdf','pdf')]:
  state=base();state['pdfs'][ident]={'type':kind,'path':path};before=set(a.out.iterdir());r=a.call(state,{'type':'convert_pdf','document_id':'invalid_'+ident,'pdf_id':ident,'letterhead_id':'plain'});assert r['error']['code']=='invalid_document' and set(a.out.iterdir())==before
 r=a.call(base(),{'type':'convert_pdf','document_id':'bad','pdf_id':'input','letterhead_id':'missing'});assert r['error']['code']=='invalid_document'
 state=base();state['assets']['sample']['path']='../outside.png';assert a.call(state,page('bad'))['error']['code']=='invalid_page'


def main():
 missing=[n for n in ['worker.py','php_launcher.py','browser_launcher.py','files.py','fixtures.py'] if not (Path('/public_tests')/n).is_file()]
 if missing:return {'harness_ok':False,'infrastructure_errors':missing,'results':[]}
 if not (ROOT/'web/index.php').is_file():return {'harness_ok':True,'results':[{'ru_id':f'RU-{i:03}','passed':False,'detail':'PHP site missing'} for i in range(1,6)]}
 rows=[]
 with tempfile.TemporaryDirectory(prefix='sweclar-docsite-') as temp:
  app=App(Path(temp))
  try:
   for i,test in enumerate([ru001,ru002,ru003,ru004,ru005],1):
    try:test(app);row={'ru_id':f'RU-{i:03}','passed':True}
    except Exception as e:row={'ru_id':f'RU-{i:03}','passed':False,'detail':str(e) or type(e).__name__,'trace':traceback.format_exc()[-1600:]}
    rows.append(row)
   if os.environ.get('DOCX_QA_OUTPUT'):shutil.copytree(app.out,os.environ['DOCX_QA_OUTPUT'],dirs_exist_ok=True)
  finally:app.close()
 return {'harness_ok':True,'results':rows}
if __name__=='__main__':
 try:r=main()
 except Exception as e:r={'harness_ok':False,'error':str(e),'results':[]}
 print(json.dumps(r))
