"""Generate synthetic local PNG, text PDFs and editable letterhead assets."""
from pathlib import Path
from PIL import Image
from pypdf import PdfWriter,PdfReader
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from docx import Document

def assets(root,lines):
 root=Path(root);root.mkdir(parents=True,exist_ok=True);Image.new('RGB',(120,80),(40,120,160)).save(root/'image.png')
 writer=PdfWriter()
 for items in [lines[:2],lines[2:]]:
  page=writer.add_blank_page(width=612,height=792);font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')});page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
  stream=DecodedStreamObject();stream.set_data(('BT /F1 12 Tf 40 750 Td '+ ' '.join(('0 -25 Td ' if i else '')+'('+text.replace(chr(92),chr(92)*2).replace('(',chr(92)+'(').replace(')',chr(92)+')')+') Tj' for i,text in enumerate(items))+' ET').encode());page[NameObject('/Contents')]=writer._add_object(stream)
 writer.write(root/'input.pdf');blank=PdfWriter();blank.add_blank_page(width=612,height=792);blank.write(root/'blank.pdf');locked=PdfWriter();locked.append(PdfReader(root/'input.pdf'));locked.encrypt('synthetic-test-password');locked.write(root/'encrypted.pdf')
 d=Document();d.sections[0].header.paragraphs[0].text='SYNTHETIC LETTERHEAD';d.add_paragraph('Sample resource note','Title');d.core_properties.author='Synthetic template author';d.core_properties.last_modified_by='Synthetic editor';d.save(root/'letterhead.docx')
