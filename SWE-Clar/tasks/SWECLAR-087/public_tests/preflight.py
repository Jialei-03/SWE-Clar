import tempfile,json
from pathlib import Path
import openpyxl
from openpyxl import Workbook,load_workbook
assert openpyxl.__version__=='3.1.5'
with tempfile.TemporaryDirectory() as temp:
 path=Path(temp)/'check.xlsx';wb=Workbook();wb.active['A1']='synthetic';wb.active['B1']='=1+2';wb.save(path);wb.close();wb=load_workbook(path);assert wb.active['A1'].value=='synthetic' and wb.active['B1'].value=='=1+2';wb.close()
print(json.dumps({'ready':True,'openpyxl':openpyxl.__version__,'actual_xlsx_roundtrip':True}))
