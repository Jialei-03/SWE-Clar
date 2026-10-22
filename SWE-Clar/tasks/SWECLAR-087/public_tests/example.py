"""Synthetic local workflow example; run from the candidate workspace."""
from pathlib import Path
import sys,json
from openpyxl import Workbook,load_workbook
root=Path.cwd();sys.path.insert(0,str(root));from inventory_sync import run_inventory_sync
work=root/'example_workspace';folder=work/'Sample';folder.mkdir(parents=True,exist_ok=True)
wb=Workbook();wb.active.title='Inventory';wb.active.append(['sku','soh','ots']);wb.active.append(['SYNTHETIC-A',None,None]);wb.save(folder/'stock.xlsx');wb.close()
(work/'sync.json').write_text(json.dumps({'customer_folders':['Sample'],'workbooks':{'Sample/stock.xlsx':{'post_process':'export','outputs':['to-order.csv']}}}))
class Inventory:
 def lookup(self,sku):return {'soh':12,'ots':4}
def post_process(path,action):
 wb=load_workbook(path);s=wb['Inventory'];(path.parent/'to-order.csv').write_text('sku,soh,ots\n'+','.join(str(c.value) for c in s[2])+'\n');wb.close()
print(json.dumps(run_inventory_sync(work,Inventory(),post_process)))
