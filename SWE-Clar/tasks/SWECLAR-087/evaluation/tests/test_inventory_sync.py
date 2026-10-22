from pathlib import Path
import builtins,importlib.util,json,os,tempfile,unittest
from unittest import mock
from openpyxl import Workbook,load_workbook
from openpyxl.styles import PatternFill,Font

import sys,subprocess,traceback
ROOT=Path(os.environ['CANDIDATE_ROOT'])


class InventoryClient:
    def __init__(self,failures=()):self.calls=[];self.failures=set(failures)
    def lookup(self,sku):
        self.calls.append(sku)
        if sku in self.failures:raise LookupError('synthetic lookup failure')
        return {'soh':len(str(sku))+10,'ots':len(str(sku))+4}


class InventorySyncTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name).resolve()
        self.config={'customer_folders':[],'workbooks':{}};self.events=[];self.callback_actions=None;self.readonly=[]
    def tearDown(self):self.temp.cleanup()
    def save_config(self):
        (self.root/'sync.json').write_text(json.dumps(self.config))
    def workbook(self,relative,rows=(),eligible=True,settings=None):
        path=self.root/relative;path.parent.mkdir(parents=True,exist_ok=True)
        customer=Path(relative).parts[0]
        if customer not in self.config['customer_folders']:self.config['customer_folders'].append(customer)
        if settings is not None:self.config['workbooks'][relative]=settings
        wb=Workbook();sheet=wb.active;sheet.title='Inventory' if eligible else 'Notes'
        sheet.append(['sku','note','qty','soh','ots','formula'])
        for row in rows:sheet.append(row)
        sheet['B2'].fill=PatternFill('solid',fgColor='AABBCC');sheet['B2'].font=Font(bold=True)
        other=wb.create_sheet('Unrelated');other['A1']='Keep this sheet';other['B1']='=SUM(1,2)'
        wb.properties.creator='Synthetic fixture';wb.save(path);wb.close();self.save_config();return path
    def cells(self,relative):
        wb=load_workbook(self.root/relative);rows=list(wb['Inventory'].values);wb.close();return rows
    def run_sync(self,client=None,callback=None):
        self.save_config();client=client or InventoryClient()
        for path in [self.root,*self.root.rglob('*')]:
            if os.getuid()==0:os.chown(path,0,0)
            path.chmod(0o777 if path.is_dir() else 0o666)
        for relative in self.readonly:(self.root/relative).chmod(0o444);(self.root/relative).parent.chmod(0o555)
        payload={'workspace':str(self.root),'failures':list(client.failures),'callbacks':self.callback_actions}
        r=subprocess.run([sys.executable,'-I','/public_tests/invoke.py',str(ROOT)],input=json.dumps(payload),capture_output=True,text=True,env={'PATH':os.environ['PATH'],'HOME':'/tmp','CLARSWE_PRIVATE_WORKER_SENTINEL':'parent-private'},timeout=20)
        self.assertEqual(r.returncode,0,r.stderr);value=json.loads(r.stdout);client.calls.extend(value['calls']);self.events=value['callbacks']
        proof=value['proof'];self.assertEqual(proof['uid'],1000);self.assertEqual(proof['cap_eff'],0);self.assertFalse(any('PRIVATE' in k for k in proof['environment_keys']))
        Path('/private/results').mkdir(parents=True,exist_ok=True)
        with open('/private/results/worker-isolation.jsonl','a') as out:out.write(json.dumps(proof)+'\n')
        return value['report']

    def test_ru_001_actual_workbook_and_report(self):
        self.workbook('A/stock.xlsx')
        report=self.run_sync()
        for field in ['processed','skipped','errors','actions','outputs','warnings']:self.assertIsInstance(report[field],list)
        self.assertEqual(report['processed'],['A/stock.xlsx'])
        self.assertEqual(report['errors'],[])

    def test_ru_002_discovery_and_output_exclusion(self):
        self.workbook('Zulu/stock.xlsx');self.workbook('Alpha/nested/stock.xlsx');self.workbook('Alpha/deep/inside/stock.xlsx')
        self.workbook('Alpha/notes.xlsx',eligible=False)
        self.workbook('run_output/should-ignore.xlsx')
        self.workbook('Beta/broken.xlsx');(self.root/'Beta/broken.xlsx').write_bytes(b'not an Excel workbook')
        (self.root/'Alpha/readme.txt').write_text('unrelated')
        missing=self.workbook('Delta/no-header.xlsx');wb=load_workbook(missing);wb['Inventory']['A1']='wrong';wb.save(missing);wb.close()
        self.config['customer_folders'].append('Alpha/nested')
        report=self.run_sync()
        self.assertEqual(report['processed'],['Alpha/deep/inside/stock.xlsx','Alpha/nested/stock.xlsx','Zulu/stock.xlsx'])
        self.assertEqual(report['skipped'],['Alpha/notes.xlsx'])
        self.assertEqual([(e['workbook'],e['stage']) for e in report['errors']],[('Beta/broken.xlsx','parse'),('Delta/no-header.xlsx','parse')])
        self.assertTrue(report['errors'][0]['error'])

    def test_ru_003_values_formulas_styles_and_blank_rows(self):
        self.workbook('A/stock.xlsx',[
            ['0012','preserve',2,-1,-1,'=C2*2'],
            [None,'blank row',3,91,92,'=C3*2'],
            ['0012','duplicate',4,-1,-1,'=C4*2'],
            ['LONG','another',5,-1,-1,'=C5*2']])
        client=InventoryClient();self.run_sync(client)
        self.assertEqual(client.calls,['0012','0012','LONG'])
        rows=self.cells('A/stock.xlsx')
        self.assertEqual(rows[1],('0012','preserve',2,14,8,'=C2*2'))
        self.assertEqual(rows[2],(None,'blank row',3,91,92,'=C3*2'))
        self.assertEqual(rows[3][3:5],(14,8))
        wb=load_workbook(self.root/'A/stock.xlsx')
        self.assertEqual(wb['Unrelated']['A1'].value,'Keep this sheet');self.assertEqual(wb['Unrelated']['B1'].value,'=SUM(1,2)')
        self.assertEqual(wb['Inventory']['B2'].fill.fgColor.rgb,'00AABBCC');self.assertTrue(wb['Inventory']['B2'].font.bold);wb.close()
        path=self.workbook('B/missing-columns.xlsx',[['NEW','unchanged']])
        wb=load_workbook(path);wb['Inventory'].delete_cols(4,3);wb.save(path);wb.close()
        self.run_sync();wb=load_workbook(path);sheet=wb['Inventory'];headers={c.value:c.column for c in sheet[1]}
        self.assertEqual(sheet.cell(2,headers['soh']).value,13);self.assertEqual(sheet.cell(2,headers['ots']).value,7);wb.close()

    def test_ru_004_real_callback_after_save_and_failure_isolation(self):
        self.workbook('A/good.xlsx',[['OK']],settings={'post_process':'recalculate'})
        self.workbook('B/partial.xlsx',[['BAD']],settings={'post_process':'must-not-run'})
        self.workbook('C/action-error.xlsx',settings={'post_process':'fail'})
        self.workbook('D/later.xlsx',settings={'post_process':'later'})
        self.workbook('E/notes.xlsx',eligible=False,settings={'post_process':'must-not-run'})
        self.callback_actions={'recalculate':{'write_file':'callback-ran.txt'},'fail':{'fail':True},'later':{'write_file':'callback-ran.txt'}}
        report=self.run_sync(InventoryClient({'BAD'}))
        self.assertEqual([(e['workbook'],e['action']) for e in self.events],[('A/good.xlsx','recalculate'),('C/action-error.xlsx','fail'),('D/later.xlsx','later')])
        self.assertEqual(self.events[0]['observed_rows'][1][3:5],[12,6])
        self.assertEqual((self.root/'A/callback-ran.txt').read_text(),'recalculate')
        self.assertEqual([(e['workbook'],e['action']) for e in report['actions']],[('A/good.xlsx','recalculate'),('D/later.xlsx','later')])
        self.assertIn(('C/action-error.xlsx','post_process'),{(e['workbook'],e['stage']) for e in report['errors']})

    def test_ru_005_collect_actual_generated_bytes_after_callback(self):
        for customer in ['A','B']:
            self.workbook(f'{customer}/stock.xlsx',[[customer]],settings={'post_process':'export','outputs':['to-order.csv']})
            (self.root/customer/'to-order.csv').write_text('stale output')
        self.workbook('C/missing.xlsx',settings={'outputs':['absent.csv']})
        self.workbook('D/lookup-failed.xlsx',[['BAD']],settings={'post_process':'export','outputs':['to-order.csv']})
        (self.root/'D/to-order.csv').write_text('must not collect')
        self.workbook('E/callback-failed.xlsx',settings={'post_process':'fail','outputs':['to-order.csv']})
        (self.root/'E/to-order.csv').write_text('must not collect')
        output=self.root/'run_output';output.mkdir();(output/'keep.csv').write_text('existing file')
        self.callback_actions={'export':{'write_file':'to-order.csv','stock_csv':True},'fail':{'fail':True}}
        report=self.run_sync(InventoryClient({'BAD'}))
        self.assertEqual(len(report['outputs']),2);self.assertEqual(len(set(report['outputs'])),2)
        self.assertEqual({(self.root/p).read_text() for p in report['outputs']},{'A,11','B,11'})
        self.assertTrue(all((self.root/p).parent==output for p in report['outputs']))
        self.assertEqual((output/'keep.csv').read_text(),'existing file')
        self.assertEqual([(w['workbook'],w['missing_output']) for w in report['warnings']],[('C/missing.xlsx','C/absent.csv')])
        previous={p:p.read_bytes() for p in output.iterdir()}
        again=self.run_sync(InventoryClient({'BAD'}))
        self.assertTrue(set(again['outputs']).isdisjoint(report['outputs']))
        self.assertTrue(all(p.read_bytes()==value for p,value in previous.items()))

    def test_ru_006_unattended_good_rows_and_later_workbooks_continue(self):
        self.workbook('A/partial.xlsx',[['BAD','keep',1,91,92],['GOOD']])
        self.workbook('B/broken.xlsx');(self.root/'B/broken.xlsx').write_bytes(b'bad archive')
        self.workbook('C/later.xlsx',[['LATER']]);client=InventoryClient({'BAD'})
        with mock.patch.object(builtins,'input',side_effect=AssertionError('interactive input is forbidden')):report=self.run_sync(client)
        self.assertEqual(client.calls,['BAD','GOOD','LATER'])
        self.assertEqual(self.cells('A/partial.xlsx')[1][3:5],(91,92));self.assertEqual(self.cells('A/partial.xlsx')[2][3:5],(14,8))
        self.assertEqual(self.cells('C/later.xlsx')[1][3:5],(15,9))
        self.assertEqual({(e['workbook'],e['stage']) for e in report['errors']},{('A/partial.xlsx','lookup'),('B/broken.xlsx','parse')})
        self.assertTrue(all(isinstance(e['error'],str) and e['error'] for e in report['errors']))
        self.workbook('D/save-failed.xlsx',[['SAVE']]);self.workbook('E/after-save.xlsx',[['AFTER']]);self.readonly=['D/save-failed.xlsx']
        report=self.run_sync(InventoryClient({'BAD'}));self.assertIn(('D/save-failed.xlsx','save'),{(e['workbook'],e['stage']) for e in report['errors']});self.assertEqual(self.cells('E/after-save.xlsx')[1][3:5],(15,9))

if __name__=='__main__':
 if not Path('/public_tests/invoke.py').is_file():print(json.dumps({'harness_ok':False,'results':[]}));sys.exit(0)
 rows=[]
 for name in sorted(n for n in dir(InventorySyncTests) if n.startswith('test_ru_')):
  suite=unittest.TestSuite([InventorySyncTests(name)]);result=unittest.TestResult();suite.run(result)
  rows.append({'ru_id':'RU-'+name.split('_')[2],'passed':result.wasSuccessful(),'detail':'\n'.join(t for _,t in result.errors+result.failures)})
 print(json.dumps({'harness_ok':True,'results':rows}))
