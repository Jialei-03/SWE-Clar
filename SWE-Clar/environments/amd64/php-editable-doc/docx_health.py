import os,json,tempfile
from pathlib import Path
import docx,lxml
from docx import Document
assert os.getuid()==1000
with tempfile.TemporaryDirectory() as directory:
 p=Path(directory)/'editable.docx';d=Document();d.sections[0].header.paragraphs[0].text='Synthetic header';d.add_paragraph('Editable paragraph');d.save(p);v=Document(p);assert v.paragraphs[0].text=='Editable paragraph' and v.sections[0].header.paragraphs[0].text=='Synthetic header';v.paragraphs[0].text='Changed';v.save(p);assert Document(p).paragraphs[0].text=='Changed'
print(json.dumps({'passed':True,'python_docx':docx.__version__,'lxml':lxml.__version__,'uid':os.getuid(),'real_editable_roundtrip':True}))
