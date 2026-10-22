import tempfile,json,os,importlib.metadata
from pathlib import Path
from openpyxl import Workbook,load_workbook
from reportlab.pdfgen import canvas
import pypdfium2
assert os.getuid()==1000
assert importlib.metadata.version('openpyxl')=='3.1.5' and importlib.metadata.version('pypdfium2')=='5.13.0'
with tempfile.TemporaryDirectory() as temp:
 root=Path(temp);wb=Workbook();wb.active['A1']='synthetic';wb.active['B1']='=1+2';wb.save(root/'test.xlsx');wb.close();wb=load_workbook(root/'test.xlsx');assert wb.active['A1'].value=='synthetic' and wb.active['B1'].value=='=1+2';wb.close()
 c=canvas.Canvas(str(root/'test.pdf'));c.drawString(60,700,'synthetic PDF');c.save();doc=pypdfium2.PdfDocument(str(root/'test.pdf'));page=doc[0];text=page.get_textpage();assert 'synthetic PDF' in text.get_text_range();image=page.render().to_pil();assert image.convert('L').getextrema()[0]<100;text.close();page.close();doc.close()
print(json.dumps({'passed':True,'uid':os.getuid(),'xlsx_roundtrip':True,'pdf_raster_text':True,'openpyxl':'3.1.5','pypdfium2':'5.13.0'}))
