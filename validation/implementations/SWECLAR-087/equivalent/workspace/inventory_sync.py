"""Local XLSX inventory synchronisation.

The public entry point in this module deliberately works with real workbooks.  It
keeps the configuration loading in :mod:`workspace` small and treats every
workbook as an independent unit so that a bad file (or a bad inventory lookup)
does not stop the rest of a run.
"""

from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
import shutil
from typing import Any

from openpyxl import load_workbook

from workspace import load_workspace


_REPORT_KEYS = ("processed", "skipped", "outputs", "warnings", "errors", "actions")


class _WorkbookValidationError(Exception):
    """Raised when an opened workbook is not an eligible inventory workbook."""


def _new_report() -> dict[str, list[Any]]:
    """Return a fresh report with the fields required by the delivery contract."""

    return {key: [] for key in _REPORT_KEYS}


def _description(error: BaseException) -> str:
    """Make an exception useful in a report even when it has an empty message."""

    text = str(error).strip()
    return text or error.__class__.__name__


def _relative_path(root: Path, path: Path) -> str:
    """Return a stable, workspace-relative path for report entries."""

    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        # Discovered workbooks are kept below root.  This fallback still gives a
        # useful value if a filesystem race or unusual symlink defeats that
        # invariant while a report is being assembled.
        try:
            return path.relative_to(root).as_posix()
        except ValueError:
            return path.as_posix()


def _normalise_relative_config_path(value: Any) -> str | None:
    """Normalise a JSON workbook key without changing its case.

    Configuration paths are documented as relative paths.  Accepting backslash
    separators as well is harmless on POSIX and makes a configuration produced
    on Windows usable when the workspace is inspected elsewhere.
    """

    if not isinstance(value, str):
        return None
    text = value.replace("\\", "/")
    while text.startswith("./"):
        text = text[2:]
    if text == ".":
        return ""
    return Path(text).as_posix()


def _is_under_run_output(root: Path, path: Path) -> bool:
    """Whether *path* is in the reserved workspace/run_output tree."""

    try:
        relative = path.resolve().relative_to(root.resolve())
    except ValueError:
        return False
    return bool(relative.parts) and relative.parts[0] == "run_output"


def _discover_workbooks(root: Path, customer_folders: Any) -> list[Path]:
    """Discover configured XLSX files in deterministic relative-path order."""

    if isinstance(customer_folders, (str, bytes)) or not customer_folders:
        folders = [customer_folders] if isinstance(customer_folders, (str, bytes)) else []
    else:
        try:
            folders = list(customer_folders)
        except TypeError:
            folders = []

    discovered: dict[Path, Path] = {}
    root_resolved = root.resolve()

    for folder_value in folders:
        if not isinstance(folder_value, (str, bytes, Path)):
            continue
        try:
            folder = root / Path(folder_value)
            folder_resolved = folder.resolve()
        except (OSError, RuntimeError, TypeError):
            continue

        # customer_folders are workspace-relative by contract.  Do not follow
        # an accidental ../ entry into an unrelated tree.
        try:
            folder_resolved.relative_to(root_resolved)
        except ValueError:
            continue

        if _is_under_run_output(root, folder_resolved):
            continue
        if not folder_resolved.exists() or not folder_resolved.is_dir():
            continue

        try:
            candidates = folder_resolved.rglob("*")
            for candidate in candidates:
                try:
                    candidate_resolved = candidate.resolve()
                except (OSError, RuntimeError):
                    continue
                if not candidate_resolved.is_file():
                    continue
                if candidate_resolved.suffix.lower() != ".xlsx":
                    continue
                try:
                    candidate_resolved.relative_to(root_resolved)
                except ValueError:
                    continue
                if _is_under_run_output(root, candidate_resolved):
                    continue
                discovered[candidate_resolved] = candidate_resolved
        except (OSError, RuntimeError):
            # An inaccessible directory can be encountered while walking a
            # customer folder.  Files that were visible before the failure are
            # still processed; there is no reliable workbook path to report for
            # a directory that could not be enumerated.
            continue

    return __import__('control_algorithms').stable_order(
        discovered.values(),
        key=lambda candidate: _relative_path(root, candidate),
    )


def _workbook_settings(config: Mapping[str, Any], relative_workbook: str) -> Mapping[str, Any]:
    """Look up optional settings for a discovered workbook."""

    configured = config.get("workbooks", {})
    if not isinstance(configured, Mapping):
        return {}

    # Prefer a direct key match, then use the normalised form for harmless path
    # spelling differences such as ``./Customer/stock.xlsx``.
    direct = configured.get(relative_workbook)
    if isinstance(direct, Mapping):
        return direct

    for key, value in configured.items():
        if _normalise_relative_config_path(key) == relative_workbook:
            return value if isinstance(value, Mapping) else {}
    return {}


def _is_nonempty_sku(value: Any) -> bool:
    """Return whether a cell contains a non-empty SKU value."""

    # A whitespace-containing cell is still a populated spreadsheet cell.  Do
    # not normalise or strip the value before handing it to the client.
    return value is not None and value != ""


def _stock_values(response: Any) -> tuple[Any, Any]:
    """Extract the documented ``soh`` and ``ots`` values from a lookup reply."""

    if isinstance(response, Mapping):
        missing = [key for key in ("soh", "ots") if key not in response]
        if missing:
            raise ValueError("lookup response is missing " + ", ".join(repr(key) for key in missing))
        return response["soh"], response["ots"]

    # Mapping is the documented shape, but accepting a small attribute-based
    # adapter makes the entry point friendly to simple test/client objects while
    # retaining the same required field names.
    try:
        return getattr(response, "soh"), getattr(response, "ots")
    except AttributeError:
        pass

    if isinstance(response, (tuple, list)) and len(response) == 2:
        return response[0], response[1]

    raise ValueError("lookup response must provide soh and ots")


def _add_stock_headers(worksheet: Any, header_columns: dict[str, int]) -> dict[str, int]:
    """Add missing exact-case stock headers after the used worksheet area."""

    next_column = worksheet.max_column + 1
    for header in ("soh", "ots"):
        if header not in header_columns:
            worksheet.cell(row=1, column=next_column).value = header
            header_columns[header] = next_column
            next_column += 1
    return header_columns


def _append_error(
    report: dict[str, list[Any]],
    workbook: str,
    stage: str,
    error: Any,
    **metadata: Any,
) -> None:
    """Append a contract-shaped error record."""

    text = str(error).strip() if error is not None else "unknown error"
    if not text:
        text = "unknown error"
    record: dict[str, Any] = {"workbook": workbook, "stage": stage, "error": text}
    record.update(metadata)
    report["errors"].append(record)


def _output_specs(settings: Mapping[str, Any]) -> list[Any]:
    value = settings.get("outputs", [])
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    if isinstance(value, (bytes, Path)):
        return [value]
    try:
        return list(value)
    except TypeError:
        return [value]


def _collision_safe_destination(run_output: Path, filename: str) -> Path:
    """Choose a destination without replacing any already-collected file."""

    destination = run_output / filename
    if not destination.exists():
        return destination

    stem = destination.stem
    suffix = destination.suffix
    number = 1
    while True:
        candidate = run_output / f"{stem}_{number}{suffix}"
        if not candidate.exists():
            return candidate
        number += 1


def _collect_outputs(
    root: Path,
    workbook_path: Path,
    relative_workbook: str,
    settings: Mapping[str, Any],
    report: dict[str, list[Any]],
) -> None:
    """Copy configured generated files after all workbook-side work is done."""

    specs = _output_specs(settings)
    if not specs:
        return

    run_output = root / "run_output"

    for spec in specs:
        if isinstance(spec, bytes):
            try:
                spec_text = spec.decode()
            except UnicodeDecodeError:
                spec_text = repr(spec)
        elif isinstance(spec, Path):
            spec_text = str(spec)
        elif isinstance(spec, str):
            spec_text = spec
        else:
            spec_text = str(spec)

        source_relative_name = spec_text.replace("\\", "/")
        source_relative = Path(source_relative_name)
        # Output specifications are relative generated filenames.  Refuse an
        # absolute path rather than allowing configuration to copy arbitrary
        # files into the run directory.
        if source_relative.is_absolute() or not source_relative_name:
            report["warnings"].append(
                {
                    "workbook": relative_workbook,
                    "missing_output": spec_text,
                    "path": spec_text,
                }
            )
            continue

        source = workbook_path.parent / source_relative
        try:
            source_resolved = source.resolve()
            source_resolved.relative_to(root.resolve())
        except (OSError, RuntimeError, ValueError):
            report["warnings"].append(
                {
                    "workbook": relative_workbook,
                    "missing_output": spec_text,
                    "path": source_relative_name,
                }
            )
            continue

        if not source_resolved.is_file():
            report["warnings"].append(
                {
                    "workbook": relative_workbook,
                    "missing_output": _relative_path(root, source_resolved),
                    "configured_output": spec_text,
                }
            )
            continue

        try:
            run_output.mkdir(parents=True, exist_ok=True)
            # Collected files are deliberately flattened into run_output.  The
            # basename is the generated filename and the suffix is collision
            # safe for workbooks that produce the same name.
            destination = _collision_safe_destination(run_output, source.name)
            shutil.copy2(source_resolved, destination)
        except Exception as error:  # output collection must not stop later files
            report["warnings"].append(
                {
                    "workbook": relative_workbook,
                    "missing_output": _relative_path(root, source_resolved),
                    "configured_output": spec_text,
                    "error": _description(error),
                }
            )
            continue

        report["outputs"].append(_relative_path(root, destination))


def _process_workbook(
    root: Path,
    workbook_path: Path,
    relative_workbook: str,
    settings: Mapping[str, Any],
    cin7_client: Any,
    post_processor: Any,
    report: dict[str, list[Any]],
) -> None:
    """Process one workbook, isolating each failure boundary."""

    workbook = None
    worksheet = None
    sku_column: int | None = None
    header_columns: dict[str, int] = {}

    try:
        workbook = load_workbook(workbook_path)
        if "Inventory" not in workbook.sheetnames:
            report["skipped"].append(relative_workbook)
            try:
                workbook.close()
            except Exception:
                pass
            workbook = None
            return

        worksheet = workbook["Inventory"]
        for cell in worksheet[1]:
            if cell.value in ("sku", "soh", "ots"):
                header_columns.setdefault(cell.value, cell.column)
        sku_column = header_columns.get("sku")
        if sku_column is None:
            raise _WorkbookValidationError(
                "Inventory worksheet does not contain an exact 'sku' header in row 1"
            )
    except _WorkbookValidationError as error:
        _append_error(report, relative_workbook, "parse", _description(error))
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
        return
    except Exception as error:
        _append_error(
            report,
            relative_workbook,
            "parse",
            f"unable to open or validate workbook: {_description(error)}",
        )
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
        return

    lookup_errors = False
    try:
        try:
            _add_stock_headers(worksheet, header_columns)
        except Exception as error:
            _append_error(
                report,
                relative_workbook,
                "save",
                f"unable to add soh/ots headers: {_description(error)}",
            )
            return

        assert sku_column is not None  # validated above
        for row_number in range(2, worksheet.max_row + 1):
            sku = worksheet.cell(row=row_number, column=sku_column).value
            if not _is_nonempty_sku(sku):
                continue

            try:
                response = cin7_client.lookup(sku)
                soh, ots = _stock_values(response)

                soh_cell = worksheet.cell(row=row_number, column=header_columns["soh"])
                ots_cell = worksheet.cell(row=row_number, column=header_columns["ots"])
                old_soh = soh_cell.value
                old_ots = ots_cell.value
                try:
                    soh_cell.value = soh
                    ots_cell.value = ots
                except Exception:
                    # A row that failed to update must retain both original
                    # stock values, just like a failed adapter lookup.
                    soh_cell.value = old_soh
                    ots_cell.value = old_ots
                    raise
            except Exception as error:
                lookup_errors = True
                _append_error(
                    report,
                    relative_workbook,
                    "lookup",
                    f"row {row_number} (SKU {sku!r}): {_description(error)}",
                    row=row_number,
                )

        try:
            workbook.save(workbook_path)
        except Exception as error:
            _append_error(
                report,
                relative_workbook,
                "save",
                f"unable to save workbook: {_description(error)}",
            )
            return
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    # A successful save is a processed workbook even when one row had an
    # isolated lookup error.  Such a partial run is intentionally not eligible
    # for callback or output collection.
    report["processed"].append(relative_workbook)
    if lookup_errors:
        return

    action = settings.get("post_process")
    callback_succeeded = True
    if action is not None and action != "" and post_processor is not None:
        try:
            # A Path is both a filesystem path and accepted by the standard
            # openpyxl/shutil APIs.  The callback therefore sees the saved file,
            # not an in-memory workbook.
            post_processor(workbook_path, action)
        except Exception as error:
            callback_succeeded = False
            _append_error(
                report,
                relative_workbook,
                "post_process",
                f"post_process action {action!r} failed: {_description(error)}",
                action=action,
            )
        else:
            report["actions"].append({"workbook": relative_workbook, "action": action})

    if callback_succeeded:
        _collect_outputs(root, workbook_path, relative_workbook, settings, report)


def run_inventory_sync(workspace, cin7_client, post_processor=None) -> dict:
    """Synchronise configured local workbooks with the supplied inventory client.

    ``workspace.load_workspace`` remains the single configuration loader.  The
    function intentionally does not execute commands from configuration: the
    optional callable is the only post-processing mechanism.
    """

    root, config = load_workspace(workspace)
    report = _new_report()

    if not isinstance(config, Mapping):
        # A valid sync.json is expected by the public contract.  Treating an
        # unusual top-level value as an empty configuration keeps an unattended
        # invocation deterministic rather than attempting attribute access on a
        # scalar.
        config = {}

    workbooks = _discover_workbooks(root, config.get("customer_folders", []))
    for workbook_path in workbooks:
        relative_workbook = _relative_path(root, workbook_path)
        settings = _workbook_settings(config, relative_workbook)
        try:
            _process_workbook(
                root,
                workbook_path,
                relative_workbook,
                settings,
                cin7_client,
                post_processor,
                report,
            )
        except Exception as error:
            # Keep an unexpected workbook-local failure from terminating an
            # unattended run.  The normal paths above use the more specific
            # parse/lookup/save/post_process stages; this is the final safety
            # boundary for an unusual workbook or filesystem implementation.
            _append_error(
                report,
                relative_workbook,
                "save",
                f"unexpected workbook processing failure: {_description(error)}",
            )

    return report